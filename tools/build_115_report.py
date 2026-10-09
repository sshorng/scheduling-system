from __future__ import annotations

from collections import OrderedDict
import os
from pathlib import Path
import re
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
import xlrd


SOURCE_XLS = Path(
    r"G:\我的雲端硬碟\工作\0教學組\01教師課務與排課\01排課與課表\01全校課表與配課總表\秘密-115建成秘密配課0820).xls"
)
TEACHER_DOC = Path(
    r"G:\我的雲端硬碟\工作\0教學組\01教師課務與排課\01排課與課表\01全校課表與配課總表\115課表-new\115學年度第1學期教師課表.docx"
)
TEMPLATE_XLSX = Path(
    r"G:\我的雲端硬碟\工作\0教學組\01教師課務與排課\01兼代課與超鐘點費用\00全校教師超鐘點與兼代課費月報核銷\114-1超鐘點暨兼課統計表(09.26) 1008.xlsx"
)
OUTPUT_XLSX = TEMPLATE_XLSX.with_name("115-1超鐘點暨兼課統計表_依配課彙整.xlsx")
TARGET_SHEET = "核章版-114-1超鐘情形表 "

SUMMARY_LABELS = {
    "實際授課時數": "actual",
    "授課時數": "actual",
    "基本授課時數": "basic",
    "標準時數": "basic",
    "減授課後時數": "reduced_after",
    "超鐘點": "source_over",
    "備註": "note",
    "授課": "native_actual",
}

PROJECT_KEYWORDS = re.compile(
    r"輔導團|領召|召集|薪傳|雙語|國樂|無人機|資優|衛星|課程小組|教師會|行政|協辦|協行|閱推|外師|ETA|共聘|校本|學困|專案"
)
LANGUAGE_TITLES = {"閩南語", "手語", "客語", "排灣族語"}
OVERTIME_OVERRIDES = {"游皓宇", "周光君", "林志鴻"}


def clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = clean(value).replace(",", "")
    match = re.fullmatch(r"[-－]?\s*\d+(?:\.\d+)?\s*(?:節)?", text)
    if not match:
        return None
    return float(text.replace("－", "-").replace("節", "").strip())


def fmt(value: float | int | None) -> str:
    if value is None:
        return ""
    value = float(value)
    if value.is_integer():
        return str(int(value))
    return f"{value:g}"


def is_class_identifier(value: object) -> bool:
    text = clean(value)
    if number(value) is not None:
        return True
    return bool(re.fullmatch(r"[7-9][A-C]", text, re.IGNORECASE))


def course_label(value: object) -> str:
    text = clean(value)
    replacements = {
        "國文科": "國文",
        "英文科": "英語",
        "數學科": "數學",
        "生物科": "生物",
        "理化科": "理化",
        "歷史科": "歷史",
        "地理科": "地理",
        "公民科": "公民",
        "音樂科": "音樂",
        "視藝科": "視覺藝術",
        "表演科": "表演藝術",
        "體育科": "體育",
        "健教科": "健康教育",
    }
    return replacements.get(text, text.removesuffix("科"))


def header_for_column(sheet: xlrd.sheet.Sheet, column: int) -> str:
    for candidate in range(column, -1, -1):
        value = sheet.cell_value(1, candidate)
        text = clean(value)
        if text and number(value) is None and text != "姓名":
            return course_label(text)
    return ""


def teacher_columns(sheet: xlrd.sheet.Sheet) -> list[int]:
    columns = []
    for column in range(1, sheet.ncols):
        name = clean(sheet.cell_value(2, column))
        if name and name != "姓名":
            columns.append(column)
    return columns


def first_summary_row(sheet: xlrd.sheet.Sheet) -> int:
    for row in range(4, sheet.nrows):
        labels = {clean(sheet.cell_value(row, c)) for c in range(sheet.ncols)}
        if labels & {
            "實際授課時數",
            "授課時數",
            "基本授課時數",
            "標準時數",
            "授課",
        }:
            return row
    return sheet.nrows


def empty_part(name: str, sheet_name: str, column: int) -> dict:
    return {
        "name": name,
        "sheet": sheet_name,
        "column": column,
        "role": "",
        "actual": 0.0,
        "has_actual": False,
        "actual_inferred": False,
        "basic": 0.0,
        "reduced_after": 0.0,
        "reduction": 0.0,
        "source_over": 0.0,
        "has_source_over": False,
        "inferred_over": 0.0,
        "has_basic": False,
        "has_reduced": False,
        "notes": [],
        "courses": [],
    }


def parse_sheet(sheet: xlrd.sheet.Sheet) -> list[dict]:
    parts = []
    summary_start = first_summary_row(sheet)
    for column in teacher_columns(sheet):
        name = clean(sheet.cell_value(2, column))
        part = empty_part(name, sheet.name, column)
        part["role"] = clean(sheet.cell_value(3, column))

        for row in range(4, sheet.nrows):
            row_labels = {clean(sheet.cell_value(row, c)) for c in range(sheet.ncols)}
            for label, key in SUMMARY_LABELS.items():
                if label not in row_labels:
                    continue
                value = sheet.cell_value(row, column)
                if key == "note":
                    text = clean(value)
                    if text and text not in part["notes"]:
                        part["notes"].append(text)
                else:
                    parsed = number(value)
                    if parsed is None:
                        continue
                    if key == "native_actual":
                        part["actual"] = parsed
                        part["has_actual"] = True
                    else:
                        part[key] = parsed
                        if key == "actual":
                            part["has_actual"] = True
                    if key == "basic":
                        part["has_basic"] = True
                    if key == "reduced_after":
                        part["has_reduced"] = True
                    if key == "source_over":
                        part["has_source_over"] = True

            if row >= summary_start:
                continue
            value = sheet.cell_value(row, column)
            hours = sheet.cell_value(row, column + 1) if column + 1 < sheet.ncols else None
            parsed_hours = number(hours)
            text = clean(value)
            if not text or parsed_hours is None or parsed_hours <= 0:
                continue
            if is_class_identifier(value):
                subject = header_for_column(sheet, column)
                class_name = text
            else:
                subject = text
                class_name = ""
            part["courses"].append(
                {
                    "subject": subject or sheet.name.strip(),
                    "class_name": class_name,
                    "hours": parsed_hours,
                }
            )
        if not part["has_actual"]:
            course_total = sum(course["hours"] for course in part["courses"])
            if course_total:
                part["actual"] = course_total
                part["has_actual"] = True
                part["actual_inferred"] = True
        if sheet.name.strip() == "本土語" and not part["has_source_over"]:
            part["inferred_over"] = part["actual"]
        if part["has_basic"] and part["has_reduced"]:
            part["reduction"] = max(0.0, part["basic"] - part["reduced_after"])
        parts.append(part)
    return parts


def merge_parts(workbook: xlrd.book.Book) -> OrderedDict[str, dict]:
    records: OrderedDict[str, dict] = OrderedDict()
    for sheet_name in workbook.sheet_names():
        for part in parse_sheet(workbook.sheet_by_name(sheet_name)):
            # The native-language sheet repeats internal teachers already listed
            # in their subject sheet; only standalone support staff are new rows.
            if part["sheet"].strip() == "本土語" and part["name"] in records:
                continue
            record = records.setdefault(
                part["name"],
                {
                    "name": part["name"],
                    "role": "",
                    "actual": 0.0,
                    "actual_inferred": False,
                    "basic": 0.0,
                    "reduced_after": 0.0,
                    "reduction": 0.0,
                    "source_over": 0.0,
                    "inferred_over": 0.0,
                    "has_basic": False,
                    "has_reduced": False,
                    "notes": [],
                    "courses": [],
                    "sheets": [],
                },
            )
            role = part["role"]
            if role and (not record["role"] or record["role"] in LANGUAGE_TITLES):
                record["role"] = role
            record["actual"] += part["actual"]
            record["actual_inferred"] = record["actual_inferred"] or part["actual_inferred"]
            if part["has_basic"]:
                record["basic"] += part["basic"]
                record["has_basic"] = True
            if part["has_reduced"]:
                record["reduced_after"] += part["reduced_after"]
                record["has_reduced"] = True
            record["reduction"] += part["reduction"]
            record["source_over"] += part["source_over"]
            record["inferred_over"] += part["inferred_over"]
            for note in part["notes"]:
                if note not in record["notes"]:
                    record["notes"].append(note)
            if part["sheet"] not in record["sheets"]:
                record["sheets"].append(part["sheet"].strip())
            record["courses"].extend(
                {
                    **course,
                    "sheet": part["sheet"].strip(),
                }
                for course in part["courses"]
            )
    return records


def teacher_order(records: OrderedDict[str, dict]) -> list[str]:
    try:
        from docx import Document

        document = Document(TEACHER_DOC)
        ordered = [
            paragraph.text.replace("教師(0827)", "").strip()
            for paragraph in document.paragraphs
            if "教師(0827)" in paragraph.text
        ]
    except Exception:
        ordered = []
    result = [name for name in ordered if name in records]
    result.extend(name for name in records if name not in result)
    return result


def class_count(value: str) -> int:
    text = (value or "").replace(" ", "").replace("，", "、")
    total = 0
    for token in text.replace("/", "、").replace("／", "、").split("、"):
        if not token:
            continue
        if (
            "-" in token
            and len(token) >= 7
            and token[:3].isdigit()
            and token[-3:].isdigit()
        ):
            total += int(token[-3:]) - int(token[:3]) + 1
        else:
            total += 1
    return total


def lesson_count(classes: str, lessons: str) -> float:
    lessons = (lessons or "").strip()
    if lessons.startswith("各"):
        return class_count(classes) * float(lessons[1:])
    try:
        return float(lessons)
    except ValueError:
        return 0.0


def teacher_doc_courses() -> dict[str, list[dict]]:
    try:
        from docx import Document

        document = Document(TEACHER_DOC)
    except Exception:
        return {}
    result = {}
    teacher_index = 0
    for paragraph in document.paragraphs:
        if "教師(0827)" not in paragraph.text:
            continue
        name = paragraph.text.replace("教師(0827)", "").strip()
        table = document.tables[2 * teacher_index + 1]
        teacher_index += 1
        entries = []
        for row in table.rows[1:]:
            cells = [cell.text.replace("\n", "／").strip() for cell in row.cells]
            for start in (0, 3):
                subject, classes, lessons = cells[start : start + 3]
                hours = lesson_count(classes, lessons)
                if not subject or hours <= 0:
                    continue
                entries.append(
                    {
                        "subject": subject,
                        "class_name": classes,
                        "hours": hours,
                        "sheet": "教師課表",
                    }
                )
        result[name] = entries
    return result


def teacher_doc_over_slots() -> dict[str, str]:
    try:
        from docx import Document

        document = Document(TEACHER_DOC)
    except Exception:
        return {}
    starts = ((3, 1), (5, 2), (7, 3), (9, 4), (12, 5), (14, 6), (16, 7), (18, 8))
    weekdays = "一二三四五"
    result = {}
    teacher_index = 0
    for paragraph in document.paragraphs:
        if "教師(0827)" not in paragraph.text:
            continue
        name = paragraph.text.replace("教師(0827)", "").strip()
        grid = document.tables[2 * teacher_index]
        teacher_index += 1
        slots = set()
        headers = [cell.text.replace("\n", "/").strip() for cell in grid.rows[0].cells[3:]]
        day_map = []
        for index, header in enumerate(headers):
            match = re.search(r"星期([一二三四五])", header)
            day_map.append(match.group(1) if match else "")
        for row_index, period in starts:
            subjects = [cell.text.replace("\n", "/").strip() for cell in grid.rows[row_index].cells[3:]]
            classes = [cell.text.replace("\n", "/").strip() for cell in grid.rows[row_index + 1].cells[3:]]
            for day, (subject, class_name) in zip(day_map, zip(subjects, classes)):
                if day and ("(超)" in subject or "(超)" in class_name):
                    slots.add((day, period))
        result[name] = "、".join(
            f"{day}{period}"
            for day, period in sorted(slots, key=lambda item: (weekdays.index(item[0]), item[1]))
        )
    return result


def teacher_doc_over_hours() -> dict[str, float]:
    try:
        from docx import Document

        document = Document(TEACHER_DOC)
    except Exception:
        return {}
    result = {}
    teacher_index = 0
    for paragraph in document.paragraphs:
        if "教師(0827)" not in paragraph.text:
            continue
        name = paragraph.text.replace("教師(0827)", "").strip()
        table = document.tables[2 * teacher_index + 1]
        teacher_index += 1
        values = [cell.text.replace("\n", "/").strip() for row in table.rows for cell in row.cells]
        for value in values:
            match = re.search(r"超鐘點\s*([-－]?\d+(?:\.\d+)?)", value)
            if match:
                result[name] = float(match.group(1).replace("－", "-"))
    return result


def apply_teacher_doc_courses(records: OrderedDict[str, dict]) -> None:
    details = teacher_doc_courses()
    for name, record in records.items():
        if details.get(name):
            record["courses"] = details[name]


def normalized_role(record: dict) -> str:
    role = record["role"].strip()
    if role in LANGUAGE_TITLES:
        return f"教支人員\n{role}"
    if role == "專任":
        return "專任教師"
    if role == "兼課":
        return "兼課教師"
    if role == "專輔":
        return "專輔教師"
    if role in {"專任教師", "共聘教師", "特教教師", "英資代理"}:
        return role
    if role.endswith("導師"):
        return role
    return f"教師兼\n{role}" if role else ""


def note_text(record: dict) -> str:
    return "\n".join(record["notes"])


def calculated_values(record: dict) -> tuple[float | None, float, float, float]:
    basic = record["basic"] if record["has_basic"] else None
    if basic == 0 and record["role"] == "兼課":
        basic = None
    reduction = record["reduction"]
    actual = record["actual"]
    expected = (basic or 0.0) - reduction
    over = actual - expected
    return basic, reduction, actual, over


def project_split(record: dict, over: float) -> tuple[float, float, str]:
    payable = max(0.0, over)
    notes = record["notes"]
    project = 0.0
    project_lines = []
    for note in notes:
        for line in re.split(r"[\r\n]+", note):
            line = line.strip()
            if not line or not PROJECT_KEYWORDS.search(line):
                continue
            project_lines.append(line)
            amounts = re.findall(r"[-－]\s*(\d+(?:\.\d+)?)", line)
            if amounts:
                project += float(amounts[0])
            else:
                project = payable
    project = min(payable, project)
    national = max(0.0, payable - project)
    label = "、".join(dict.fromkeys(project_lines))
    return project, national, label


def course_summary(record: dict) -> tuple[str, str, str]:
    groups: OrderedDict[str, dict] = OrderedDict()
    for course in record["courses"]:
        subject = course["subject"] or "未標示課程"
        group = groups.setdefault(subject, {"hours": 0.0, "classes": [], "sheets": []})
        group["hours"] += course["hours"]
        if course["class_name"] and course["class_name"] not in group["classes"]:
            group["classes"].append(course["class_name"])
        if course["sheet"] not in group["sheets"]:
            group["sheets"].append(course["sheet"])
    subjects = "、".join(groups)
    class_parts = []
    for subject, group in groups.items():
        if group["classes"]:
            class_parts.append(f"{subject}：{'、'.join(group['classes'])}")
        else:
            class_parts.append(f"{subject}（{fmt(group['hours'])}節）")
    classes = "；".join(class_parts)
    sources = "、".join(record["sheets"])
    remarks = f"配課來源：{sources}" if sources else ""
    if "本土語" in record["sheets"]:
        remarks += "；本土語支援納入"
    return subjects, classes, remarks


def raw_source_over(record: dict, over_hours: dict[str, float] | None = None) -> float:
    if over_hours and record["name"] in OVERTIME_OVERRIDES and record["name"] in over_hours:
        return over_hours[record["name"]]
    return record["source_over"] + record["inferred_over"]


def effective_source_over(record: dict, over_hours: dict[str, float] | None = None) -> float:
    # Negative source values indicate underload, not payable over-hours.
    return max(0.0, raw_source_over(record, over_hours))


def set_cell_value(sheet, row: int, column: int, value: object) -> None:
    sheet.cell(row=row, column=column).value = value


def clear_cells(sheet, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
    for row in range(start_row, end_row + 1):
        for column in range(start_col, end_col + 1):
            sheet.cell(row=row, column=column).value = None


def cache_formula_values(output: Path, values: dict[str, float]) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", namespace)
    with zipfile.ZipFile(output, "r") as source_zip:
        entries = {name: source_zip.read(name) for name in source_zip.namelist()}
    sheet_name = "xl/worksheets/sheet1.xml"
    root = ET.fromstring(entries[sheet_name])
    for cell in root.findall(f".//{{{namespace}}}c"):
        reference = cell.get("r")
        if reference not in values:
            continue
        value_node = cell.find(f"{{{namespace}}}v")
        if value_node is None:
            value_node = ET.SubElement(cell, f"{{{namespace}}}v")
        value = values[reference]
        value_node.text = fmt(value)
    entries[sheet_name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(temporary_path, "w", zipfile.ZIP_DEFLATED) as destination_zip:
            for name, content in entries.items():
                destination_zip.writestr(name, content)
        os.replace(temporary_path, output)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def write_report(records: OrderedDict[str, dict], output: Path) -> None:
    workbook = openpyxl.load_workbook(TEMPLATE_XLSX)
    sheet = workbook[TARGET_SHEET]
    names = teacher_order(records)
    over_slots = teacher_doc_over_slots()
    over_hours = teacher_doc_over_hours()
    data_start = 3
    data_end = data_start + len(names) - 1
    template_data_end = 78
    total_row = 79
    formula_cache: dict[str, float] = {}
    total_d = total_e = total_f = total_g = total_n = total_o = total_p = 0.0

    sheet["A1"] = "臺北市立建成國民中學115學年度第一學期教師超鐘點暨兼課統計表"
    clear_cells(sheet, data_start, template_data_end, 1, 16)

    for index, name in enumerate(names, start=1):
        row = data_start + index - 1
        record = records[name]
        basic, reduction, actual, over = calculated_values(record)
        reduction_note = note_text(record)

        source_over = effective_source_over(record, over_hours)
        if source_over > 0:
            subjects, classes, _ = course_summary(record)
        else:
            subjects, classes = "", ""
        use_source_over = abs(source_over - over) > 0.01
        if use_source_over:
            g_value: object = source_over
        else:
            g_value = f"=F{row}-(D{row}-E{row})"
        project, national, project_label = project_split(record, source_over)
        total_d += basic or 0.0
        total_e += reduction
        total_f += actual
        total_g += source_over
        total_n += project
        total_o += national
        total_p += project + national
        if isinstance(g_value, str) and g_value.startswith("="):
            formula_cache[f"G{row}"] = over
        formula_cache[f"P{row}"] = project + national

        values = {
            1: index,
            2: normalized_role(record),
            3: name,
            4: basic,
            5: reduction,
            6: actual,
            7: g_value,
            8: reduction_note or None,
            9: subjects or None,
            10: classes or None,
            11: over_slots.get(name) or None,
            12: (f"{project_label}（{fmt(project)}）" if project > 0 else None),
            13: (f"國教{fmt(national)}" if national > 0 else None),
            14: project,
            15: national,
            16: f"=SUM(N{row}:O{row})",
        }
        for column, value in values.items():
            set_cell_value(sheet, row, column, value)

    clear_cells(sheet, data_end + 1, template_data_end, 1, 16)

    for column, letter in ((4, "D"), (5, "E"), (6, "F"), (7, "G"), (14, "N"), (15, "O"), (16, "P")):
        set_cell_value(sheet, total_row, column, f"=SUM({letter}{data_start}:{letter}{data_end})")
    formula_cache.update(
        {
            f"D{total_row}": total_d,
            f"E{total_row}": total_e,
            f"F{total_row}": total_f,
            f"G{total_row}": total_g,
            f"N{total_row}": total_n,
            f"O{total_row}": total_o,
            f"P{total_row}": total_p,
        }
    )
    set_cell_value(
        sheet,
        80,
        6,
        "資料來源：秘密-115建成秘密配課0820；D/E/F取各科配課彙整，G以來源超鐘點與公式核對；本土語支援納入。",
    )

    clear_cells(sheet, 81, 90, 13, 15)
    set_cell_value(sheet, 81, 13, "分項合計")
    set_cell_value(sheet, 81, 14, f"=SUM(N{data_start}:N{data_end})")
    formula_cache["N81"] = total_n
    set_cell_value(sheet, 82, 13, "國教合計")
    set_cell_value(sheet, 82, 15, f"=SUM(O{data_start}:O{data_end})")
    formula_cache["O82"] = total_o
    set_cell_value(sheet, 83, 13, "支付小計")
    set_cell_value(sheet, 83, 15, f"=SUM(P{data_start}:P{data_end})")
    formula_cache["O83"] = total_p

    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except AttributeError:
        pass
    workbook.save(output)
    cache_formula_values(output, formula_cache)
    print(f"created: {output}")
    print(f"records: {len(names)}")
    print(f"data rows: {data_start}:{data_end}")


def inspect(records: OrderedDict[str, dict]) -> None:
    over_hours = teacher_doc_over_hours()
    for name in teacher_order(records):
        record = records[name]
        basic, reduction, actual, over = calculated_values(record)
        source_over = effective_source_over(record, over_hours)
        raw_over = raw_source_over(record)
        print(
            f"{name}\trole={normalized_role(record).replace(chr(10), '/') }\t"
            f"D={fmt(basic)} E={fmt(reduction)} F={fmt(actual)} G={fmt(over)}\t"
            f"sourceG={fmt(source_over)} rawG={fmt(raw_over)}\t"
            f"notes={'；'.join(record['notes'])}"
        )


def main() -> None:
    if not SOURCE_XLS.exists():
        raise FileNotFoundError(SOURCE_XLS)
    if not TEMPLATE_XLSX.exists():
        raise FileNotFoundError(TEMPLATE_XLSX)
    source = xlrd.open_workbook(SOURCE_XLS, formatting_info=False)
    records = merge_parts(source)
    apply_teacher_doc_courses(records)
    if "--inspect" in sys.argv:
        inspect(records)
    else:
        write_report(records, OUTPUT_XLSX)


if __name__ == "__main__":
    main()
